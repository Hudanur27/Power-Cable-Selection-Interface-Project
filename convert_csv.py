
import pandas as pd
from openpyxl import load_workbook

def get_rgb(cell):
    fill = cell.fill
    if fill and fill.fill_type == 'solid' and fill.fgColor.type == 'rgb':
        rgb = fill.fgColor.rgb
        if rgb and len(rgb) == 8:  # Includes alpha
            r, g, b = int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16)
            return r, g, b
    return None

def is_grayish(rgb):
    if rgb is None:
        return False
    r, g, b = rgb
    return abs(r - g) < 10 and abs(r - b) < 10 and abs(g - b) < 10

def excel_to_clean_csv(excel_path, sheet_name=0, output_csv='output.csv'):
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[sheet_name]

    data = []
    for row in ws.iter_rows():
        row_data = []
        for cell in row:
            rgb = get_rgb(cell)
            if is_grayish(rgb):
                row_data.append("NA")
            else:
                row_data.append(cell.value if cell.value is not None else "")
        data.append(row_data)

    df = pd.DataFrame(data)
    df.to_csv(output_csv, index=False, header=False)
    print(f"Saved cleaned CSV to {output_csv}")




# Example usage
excel_to_clean_csv("C:/Users/Casper/Desktop/EE374 Project Cable List.xlsx")

